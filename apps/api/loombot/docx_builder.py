"""Build Word documents from simple Markdown, and turn Word documents back into it.

Files and Ask write DOCX files as Markdown: ``#`` headings, ``-`` and ``1.``
lists, ``>`` quotes, ``|`` tables, ``---`` rules, and inline ``**bold**``,
``*italic*``, ``code``, and ``[links](url)``. The reverse conversion lets a
DOCX be edited as text in the Files tab and by Ask.
"""

from __future__ import annotations

import io
import re
from typing import Any

_HEADING = re.compile(r"^\s{0,3}(#{1,6})\s+(.*?)(?:\s+#+)?\s*$")
_BULLET = re.compile(r"^(\s*)[-*+]\s+(.*)$")
_NUMBERED = re.compile(r"^(\s*)\d{1,9}[.)]\s+(.*)$")
_QUOTE = re.compile(r"^\s{0,3}>\s?(.*)$")
_RULE = re.compile(r"^\s{0,3}([-*_])(?:\s*\1){2,}\s*$")
_TABLE_SEPARATOR = re.compile(r"^\s*\|?\s*:?-{1,}:?\s*(\|\s*:?-{1,}:?\s*)*\|?\s*$")
_INLINE = re.compile(
    r"(\*\*(?P<bold>.+?)\*\*|__(?P<bold2>.+?)__|`(?P<code>[^`]+)`"
    r"|\[(?P<label>[^\]]+)\]\((?P<url>[^)\s]+)\)"
    r"|(?<![\w*])\*(?P<italic>[^*\n]+?)\*(?![\w*])|(?<![\w_])_(?P<italic2>[^_\n]+?)_(?![\w_]))"
)


def _cells(line: str) -> list[str]:
    row = line.strip()
    if row.startswith("|"):
        row = row[1:]
    if row.endswith("|"):
        row = row[:-1]
    return [cell.strip() for cell in row.split("|")]


def _add_runs(paragraph: Any, text: str, *, bold: bool = False) -> None:
    """Add ``text`` to ``paragraph``, turning inline Markdown into run formatting."""

    position = 0
    for match in _INLINE.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position:match.start()]).bold = bold or None
        groups = match.groupdict()
        if groups["bold"] or groups["bold2"]:
            run = paragraph.add_run(groups["bold"] or groups["bold2"])
            run.bold = True
        elif groups["code"]:
            run = paragraph.add_run(groups["code"])
            run.font.name = "Consolas"
            run.bold = bold or None
        elif groups["label"]:
            run = paragraph.add_run(f"{groups['label']} ({groups['url']})")
            run.bold = bold or None
        else:
            run = paragraph.add_run(groups["italic"] or groups["italic2"])
            run.italic = True
            run.bold = bold or None
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:]).bold = bold or None


def _list_style(document: Any, base: str, depth: int) -> str:
    name = base if depth == 0 else f"{base} {min(depth + 1, 3)}"
    try:
        document.styles[name]
        return name
    except KeyError:
        return base


def markdown_to_docx(markdown: str, *, title: str | None = None) -> bytes:
    """A .docx file built from Markdown text."""

    import docx

    document = docx.Document()
    if title:
        document.core_properties.title = title[:250]
    lines = (markdown or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")
    paragraph: Any = None  # the open plain paragraph, joined across soft line breaks
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            paragraph = None
            i += 1
            continue
        heading = _HEADING.match(line)
        if heading:
            level = len(heading.group(1))
            _add_runs(document.add_heading("", level=min(level, 9)), heading.group(2))
            paragraph = None
            i += 1
            continue
        if _RULE.match(line):
            document.add_paragraph()
            paragraph = None
            i += 1
            continue
        if stripped.startswith("|") and "|" in stripped[1:]:
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not _TABLE_SEPARATOR.match(lines[i]):
                    rows.append(_cells(lines[i]))
                i += 1
            if rows:
                width = max(len(row) for row in rows)
                table = document.add_table(rows=len(rows), cols=width)
                try:
                    table.style = "Table Grid"
                except (KeyError, ValueError):
                    pass
                for r, row in enumerate(rows):
                    for c in range(width):
                        cell = table.cell(r, c)
                        cell.text = ""
                        _add_runs(cell.paragraphs[0], row[c] if c < len(row) else "", bold=r == 0)
            paragraph = None
            continue
        bullet = _BULLET.match(line)
        numbered = None if bullet else _NUMBERED.match(line)
        if bullet or numbered:
            match = bullet or numbered
            depth = len(match.group(1).replace("\t", "    ")) // 2
            style = _list_style(document, "List Bullet" if bullet else "List Number", depth)
            _add_runs(document.add_paragraph(style=style), match.group(2))
            paragraph = None
            i += 1
            continue
        quote = _QUOTE.match(line)
        if quote:
            try:
                target = document.add_paragraph(style="Quote")
            except KeyError:
                target = document.add_paragraph()
            _add_runs(target, quote.group(1))
            paragraph = None
            i += 1
            continue
        if paragraph is None:
            paragraph = document.add_paragraph()
        else:
            paragraph.add_run().add_break()
        _add_runs(paragraph, stripped)
        i += 1
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _wrap(text: str, marker: str) -> str:
    """``**text**`` with surrounding spaces kept outside the markers."""

    stripped = text.strip()
    if not stripped:
        return text
    lead = text[: len(text) - len(text.lstrip())]
    trail = text[len(text.rstrip()):]
    return f"{lead}{marker}{stripped}{marker}{trail}"


def paragraph_markdown(paragraph: Any) -> str:
    """A Word paragraph's text with bold, italic, code, and links as inline Markdown."""

    from docx.text.hyperlink import Hyperlink

    parts: list[str] = []
    try:
        items = list(paragraph.iter_inner_content())
    except AttributeError:  # older python-docx
        items = list(paragraph.runs)
    for item in items:
        if isinstance(item, Hyperlink):
            label = item.text
            address = getattr(item, "address", "") or ""
            parts.append(f"[{label}]({address})" if address and label else label)
            continue
        text = item.text
        if not text:
            continue
        if getattr(item.font, "name", None) == "Consolas":
            text = f"`{text}`"
        elif item.bold and item.italic:
            text = _wrap(text, "***")
        elif item.bold:
            text = _wrap(text, "**")
        elif item.italic:
            text = _wrap(text, "*")
        parts.append(text)
    # Adjacent runs with the same formatting ("**a****b**") read better merged.
    return "".join(parts).replace("****", "").strip()


def blocks_to_markdown(blocks: list[dict[str, Any]]) -> str:
    """Markdown for the heading/paragraph/list/table blocks of a Word document."""

    out: list[str] = []
    number = 0
    for block in blocks:
        kind = block.get("type")
        number = number + 1 if kind == "list_item" and block.get("ordered") else 0
        text = block.get("markdown") or block.get("text") or ""
        if kind == "heading":
            out.append(f"{'#' * min(int(block.get('level') or 1), 6)} {block['text']}")
        elif kind == "list_item":
            out.append(f"{number}. {text}" if block.get("ordered") else f"- {text}")
        elif kind == "table":
            rows = [[" ".join(str(cell).split()) for cell in row] for row in block.get("rows") or []]
            if not rows:
                continue
            width = max(len(row) for row in rows)
            rows = [row + [""] * (width - len(row)) for row in rows]
            table = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join(["---"] * width) + " |"]
            table += ["| " + " | ".join(row) + " |" for row in rows[1:]]
            out.append("\n".join(table))
        else:
            out.append(str(text))
    # Consecutive list items stay together; everything else is its own paragraph.
    def is_item(part: str) -> bool:
        return part.startswith("- ") or bool(re.match(r"\d+\. ", part))

    text = ""
    for index, part in enumerate(out):
        if index:
            text += "\n" if is_item(part) and is_item(out[index - 1]) else "\n\n"
        text += part
    return text
