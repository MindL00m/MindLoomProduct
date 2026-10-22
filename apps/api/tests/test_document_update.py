"""Updated document versions: segmenting, applying edits per format, and the full flow."""

from __future__ import annotations

import io
from datetime import datetime, timezone
from types import SimpleNamespace

import pytest

from loombot import document_update as du
from loombot import files as files_service


def _docx() -> bytes:
    import docx

    document = docx.Document()
    document.add_heading("Q3 Status Report", level=1)
    paragraph = document.add_paragraph("Revenue was ")
    run = paragraph.add_run("$1.2M")
    run.bold = True
    paragraph.add_run(" for the quarter.")
    document.add_paragraph("Launch date: September 30.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Project", "Status"
    table.cell(1, 0).text, table.cell(1, 1).text = "Tiramisu", "In progress"
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def _xlsx() -> bytes:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Prices"
    sheet.append(["Item", "Price"])
    sheet.append(["Flour", 4.0])
    sheet.append(["Sugar", 3.0])
    sheet["B4"] = "=SUM(B2:B3)"
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def _pdf() -> bytes:
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    page.insert_textbox(fitz.Rect(72, 72, 520, 110), "Headcount: 42 engineers", fontsize=14)
    page.insert_textbox(fitz.Rect(72, 130, 520, 170), "Office: Pittsburgh", fontsize=11)
    data = doc.tobytes()
    doc.close()
    return data


def test_docx_edits_keep_formatting_and_structure() -> None:
    data = _docx()
    plan = du.segment_document("docx", data, "")
    texts = {s.id: s.text for s in plan.segments if s.editable}
    revenue_id = next(i for i, t in texts.items() if t.startswith("Revenue"))
    status_id = next(i for i, t in texts.items() if t == "In progress")
    assert any(s.role.startswith("table cell") for s in plan.segments)

    updated = du.apply_edits(
        "docx", data, "",
        {revenue_id: "Revenue was $1.5M for the quarter.", status_id: "Launched"},
        [(revenue_id, "Margin improved to 31%.")],
    )
    import docx

    document = docx.Document(io.BytesIO(updated))
    body = [p.text for p in document.paragraphs]
    assert body[:4] == ["Q3 Status Report", "Revenue was $1.5M for the quarter.",
                        "Margin improved to 31%.", "Launch date: September 30."]
    assert document.paragraphs[1].runs[0].bold is None  # first run's formatting kept
    assert document.tables[0].cell(1, 1).text == "Launched"
    assert document.paragraphs[0].style.name == "Heading 1"


def test_xlsx_edits_keep_numbers_and_formulas() -> None:
    data = _xlsx()
    plan = du.segment_document("xlsx", data, "")
    by_id = {s.id: s for s in plan.segments}
    assert by_id["Prices!B4"].editable is False  # formulas are read-only
    updated = du.apply_edits("xlsx", data, "", {"Prices!B2": "4.50"}, [])
    import openpyxl

    sheet = openpyxl.load_workbook(io.BytesIO(updated))["Prices"]
    assert sheet["B2"].value == 4.5 and sheet["B4"].value == "=SUM(B2:B3)"


def test_pdf_edits_replace_text_in_place() -> None:
    data = _pdf()
    plan = du.segment_document("pdf", data, "")
    target = next(s for s in plan.segments if "Headcount" in s.text)
    updated = du.apply_edits("pdf", data, "", {target.id: "Headcount: 57 engineers"}, [])
    text = files_service.pdf_text(updated)
    assert "57 engineers" in text and "42 engineers" not in text and "Pittsburgh" in text


def test_csv_and_txt_edits_and_inserts() -> None:
    content = "item,price\nflour,4.00\nsugar,3.00\n"
    out = du.apply_edits("csv", b"", content, {"R2C2": "4.50"}, [("R3", "salt,1.00"), ("R3", "yeast,2.00")])
    assert out.splitlines() == ["item,price", "flour,4.50", "sugar,3.00", "salt,1.00", "yeast,2.00"]
    txt = "Status: planning\n\nOwner: Ada\n"
    out = du.apply_edits("txt", b"", txt, {"t0": "Status: shipped"}, [("t1", "Next: retro on Friday")])
    assert out == "Status: shipped\n\nOwner: Ada\n\nNext: retro on Friday\n"


async def test_full_flow_cites_evidence_and_saves_a_new_version(monkeypatch) -> None:
    data = _docx()
    chunk = SimpleNamespace(
        chunk_id="c1", summary="Q3 revenue closed at $1.5M", raw_text="Finance: Q3 revenue was $1.5M.",
        end_time=datetime(2026, 10, 1, tzinfo=timezone.utc), retrieval_score=0.9,
        citation=SimpleNamespace(source_label="Finance update"),
    )

    async def fake_bytes(org_id, user_id, file_id):
        return data, {"name": "Q3 report.docx", "kind": "docx", "file_id": file_id}

    async def fake_retrieve(query, history, org_id, tokens):
        return SimpleNamespace(chunks=[chunk])

    async def tokens(org_id, user_id):
        return ["user:u"]

    async def no_cite(chunks, org_id):
        return None

    plan = du.segment_document("docx", data, "")
    revenue_id = next(s.id for s in plan.segments if s.text.startswith("Revenue"))
    launch_id = next(s.id for s in plan.segments if s.text.startswith("Launch"))
    replies = [
        {"focus": "Q3 report", "queries": ["Q3 revenue"]},
        {"summary": "Revenue updated.", "edits": [
            {"id": revenue_id, "text": "Revenue was $1.5M for the quarter.", "chunk_ids": ["c1"]},
            {"id": launch_id, "text": "Launch date: November 3.", "chunk_ids": []},  # uncited
            {"id": "p999", "text": "nope", "chunk_ids": ["c1"]},                       # unknown
        ], "inserts": []},
    ]

    async def fake_chat(org_id, system, user):
        return replies.pop(0)

    saved: dict = {}

    async def fake_create(org_id, user_id, name, kind, payload, *, edited_by="user", based_on=None):
        saved.update(name=name, kind=kind, payload=payload, edited_by=edited_by, based_on=based_on)
        return {"file_id": "new", "name": name, "kind": kind, "content": ""}

    async def fake_list(org_id, user_id):
        return [{"name": "Q3 report.docx"}, {"name": "Q3 report (updated 2026-10-22).docx"}]

    import brain.retrieval
    import keeper.auth
    import loombot.answerer
    from lens import status_trackers

    monkeypatch.setattr(files_service, "get_file_bytes", fake_bytes)
    monkeypatch.setattr(files_service, "create_binary_file", fake_create)
    monkeypatch.setattr(files_service, "list_files", fake_list)
    monkeypatch.setattr(brain.retrieval, "retrieve", fake_retrieve)
    monkeypatch.setattr(keeper.auth, "get_user_access_tokens", tokens)
    monkeypatch.setattr(loombot.answerer, "_attach_citations", no_cite)
    monkeypatch.setattr(status_trackers, "_chat_json", fake_chat)

    result = await du.create_updated_version("o", "u", "f1")

    assert result["changed"] is True and result["based_on"] == "Q3 report.docx"
    assert [c["after"] for c in result["changes"]] == ["Revenue was $1.5M for the quarter."]
    assert result["changes"][0]["sources"] == ["Finance update (2026-10-01)"]
    assert saved["kind"] == "docx" and saved["based_on"] == "f1" and saved["edited_by"] == "loombot"
    assert saved["name"].startswith("Q3 report (updated ") and saved["name"].endswith(".docx")
    import docx

    paragraphs = [p.text for p in docx.Document(io.BytesIO(saved["payload"])).paragraphs]
    assert "Revenue was $1.5M for the quarter." in paragraphs
    assert "Launch date: September 30." in paragraphs  # uncited edit dropped


def test_unique_name_avoids_collisions() -> None:
    existing = {"Report (updated 2026-10-22).docx"}
    assert files_service.unique_name(existing, "report (updated 2026-10-22).docx") == "report (updated 2026-10-22) (2).docx"
    assert files_service.unique_name(existing, "Other.docx") == "Other.docx"


@pytest.mark.parametrize("name,kind", [("a.csv", "csv"), ("b.TXT", "txt"), ("c.docx", "docx"), ("d.pdf", "pdf"), ("e.xlsx", "xlsx"), ("f.png", None)])
def test_kind_from_name(name, kind) -> None:
    assert files_service.kind_from_name(name) == kind
