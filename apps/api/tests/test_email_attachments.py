"""Email attachments are read in live search and indexed during mailbox sync."""

from __future__ import annotations

import base64
import io
import json
from datetime import datetime, timezone
from types import SimpleNamespace

import fitz  # PyMuPDF
import httpx
import pytest

from loombot import ask_agent, data_tools
from loombot import files as files_service
from weaver import email_content, google_workspace


def _b64(data: bytes | str) -> str:
    raw = data.encode() if isinstance(data, str) else data
    return base64.urlsafe_b64encode(raw).decode().rstrip("=")


def _pdf(text: str) -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), text, fontsize=11)
    data = doc.tobytes()
    doc.close()
    return data


def _scan_pdf(pages: int = 2) -> bytes:
    """Pages that are only a picture of text, like a scanner produces."""

    import fitz

    picture = fitz.open()
    source = picture.new_page()
    source.insert_text((50, 60), "Receipt total 42.00", fontsize=14)
    png = source.get_pixmap(matrix=fitz.Matrix(2, 2)).tobytes("png")
    doc = fitz.open()
    for _ in range(pages):
        doc.new_page().insert_image(fitz.Rect(0, 0, 595, 842), stream=png)
    return doc.tobytes()


def _blank_pdf(pages: int = 2) -> bytes:
    doc = fitz.open()
    for _ in range(pages):
        doc.new_page()
    data = doc.tobytes()
    doc.close()
    return data


def _docx(text: str) -> bytes:
    import docx

    document = docx.Document()
    document.add_heading("Contract", level=1)
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


INVOICE = _pdf("Invoice INV-2041 total due 1,250.00 USD by 30 Nov 2026")

MESSAGE = {
    "id": "m1",
    "threadId": "t1",
    "internalDate": str(int(datetime(2026, 10, 21, tzinfo=timezone.utc).timestamp() * 1000)),
    "snippet": "Please find the invoice attached",
    "labelIds": ["INBOX"],
    "payload": {
        "mimeType": "multipart/mixed",
        "headers": [
            {"name": "Subject", "value": "October invoice"},
            {"name": "From", "value": "Acme Billing <billing@acme.com>"},
            {"name": "To", "value": "me@acme.com"},
        ],
        "parts": [
            {"mimeType": "text/plain", "filename": "", "body": {"data": _b64("Please find the invoice attached.")}},
            {
                "mimeType": "application/pdf", "filename": "INV-2041.pdf",
                "headers": [{"name": "Content-Disposition", "value": 'attachment; filename="INV-2041.pdf"'}],
                "body": {"attachmentId": "att-pdf", "size": len(INVOICE)},
            },
            {
                "mimeType": "image/png", "filename": "logo.png",
                "headers": [
                    {"name": "Content-Disposition", "value": "inline"},
                    {"name": "Content-ID", "value": "<logo>"},
                ],
                "body": {"attachmentId": "att-logo", "size": 2048},
            },
            {
                "mimeType": "application/msword", "filename": "old.doc",
                "body": {"attachmentId": "att-doc", "size": 4000},
            },
        ],
    },
}


def test_gmail_attachment_refs_flags_inline_images() -> None:
    refs = email_content.gmail_attachment_refs(MESSAGE["payload"])
    assert [(r.filename, r.inline) for r in refs] == [
        ("INV-2041.pdf", False), ("logo.png", True), ("old.doc", False),
    ]
    assert email_content.skip_reason(refs[0]) is None
    assert "logo" in email_content.skip_reason(refs[1])
    assert "Older Office" in email_content.skip_reason(refs[2])


async def test_extract_attachment_reads_common_formats() -> None:
    pdf = await email_content.extract_attachment("INV-2041.pdf", "application/pdf", INVOICE)
    assert pdf.method == "text" and pdf.pages == 1
    assert "INV-2041" in pdf.text and "1,250.00" in pdf.text

    word = await email_content.extract_attachment(
        "contract.docx", "application/octet-stream", _docx("Renewal term is 24 months.")
    )
    assert "Renewal term is 24 months." in word.text

    sheet = await email_content.extract_attachment("prices.csv", "text/csv", b"item,price\nflour,4.50\n")
    assert "flour,4.50" in sheet.text

    html = await email_content.extract_attachment(
        "receipt.html", "text/html", b"<html><style>x{}</style><p>Order <b>#88</b></p></html>"
    )
    assert html.text == "Order #88"

    forwarded = (
        b"Subject: Signed SOW\r\nFrom: Priya <priya@acme.com>\r\nDate: Tue, 20 Oct 2026 10:00:00 +0000\r\n"
        b"Content-Type: text/plain\r\n\r\nThe SOW is signed and starts Monday.\r\n"
    )
    email_text = await email_content.extract_attachment("fwd.eml", "message/rfc822", forwarded)
    assert "Signed SOW" in email_text.text and "starts Monday" in email_text.text


async def test_scanned_pdf_and_images_are_flagged_without_ocr(monkeypatch) -> None:
    monkeypatch.setattr(email_content, "_ocr_enabled", lambda: False)
    scanned = await email_content.extract_attachment("scan.pdf", "application/pdf", _scan_pdf())
    assert scanned.text == "" and "scanned" in scanned.note.lower()
    photo = await email_content.extract_attachment("receipt.jpg", "image/jpeg", b"x" * 40_000)
    assert photo.text == "" and "OCR" in photo.note


async def test_scanned_pdf_is_transcribed_with_ocr(monkeypatch) -> None:
    seen: list[int] = []

    async def fake_ocr(images, *, label):
        seen.append(len(images))
        return "Receipt total 42.00"

    monkeypatch.setattr(email_content, "_ocr_enabled", lambda: True)
    monkeypatch.setattr(email_content, "ocr_images", fake_ocr)
    scanned = await email_content.extract_attachment("scan.pdf", "application/pdf", _scan_pdf(2))
    assert scanned.method == "ocr"
    assert scanned.text == "[Page 1]\nReceipt total 42.00\n\n[Page 2]\nReceipt total 42.00"
    assert seen == [1, 1]  # one page per OCR call, so text stays with its page


def test_split_for_chunks_keeps_pieces_small() -> None:
    text = "\n\n".join(f"Paragraph {i} " + "word " * 120 for i in range(12))
    pieces = email_content.split_for_chunks(text, max_tokens=300)
    from brain.chunker import _ENCODING

    assert len(pieces) > 4
    assert all(len(_ENCODING.encode(piece)) <= 300 for piece in pieces)
    assert "Paragraph 0" in pieces[0] and "Paragraph 11" in pieces[-1]


def _connect_gmail(monkeypatch) -> None:
    from weaver import connection_setup, integrations

    async def fake_connection(org_id, user_id, provider):
        return SimpleNamespace(provider=provider)

    async def fake_policy(org_id, user_id, provider):
        return SimpleNamespace(included_resources=json.dumps(["mailbox:gmail"]))

    async def fake_token(org_id, user_id):
        return "google-token", "me@acme.com"

    monkeypatch.setattr(integrations, "_get_connection", fake_connection)
    monkeypatch.setattr(connection_setup, "get_policy", fake_policy)
    monkeypatch.setattr(google_workspace, "_workspace_access_token", fake_token)


def _handler(request: httpx.Request) -> httpx.Response:
    path = request.url.path
    if path.endswith("/messages"):
        return httpx.Response(200, json={"messages": [{"id": "m1"}], "resultSizeEstimate": 1})
    if path.endswith("/attachments/att-pdf"):
        return httpx.Response(200, json={"data": _b64(INVOICE), "size": len(INVOICE)})
    if "/attachments/" in path:
        raise AssertionError(f"skipped attachments must not be downloaded: {path}")
    return httpx.Response(200, json=MESSAGE)


def _mock_http(monkeypatch) -> list[httpx.Request]:
    requests: list[httpx.Request] = []
    real_client = httpx.AsyncClient

    def recording(request: httpx.Request) -> httpx.Response:
        requests.append(request)
        return _handler(request)

    monkeypatch.setattr(
        data_tools.httpx, "AsyncClient",
        lambda **kwargs: real_client(transport=httpx.MockTransport(recording), **kwargs),
    )
    return requests


async def test_search_emails_returns_attachment_text(monkeypatch) -> None:
    _connect_gmail(monkeypatch)
    _mock_http(monkeypatch)

    result = await data_tools.search_emails("org-1", "u-1", sender="acme")

    message = result["messages"][0]
    assert message["body"] == "Please find the invoice attached."
    by_name = {item["filename"]: item for item in message["attachments"]}
    assert by_name["INV-2041.pdf"]["read"] is True
    assert "1,250.00" in by_name["INV-2041.pdf"]["text"]
    assert by_name["logo.png"]["read"] is False
    assert by_name["old.doc"]["read"] is False and "Older Office" in by_name["old.doc"]["note"]
    assert "attachments" in result["note"]


async def test_read_email_attachment_returns_full_text(monkeypatch) -> None:
    _connect_gmail(monkeypatch)
    _mock_http(monkeypatch)

    result = await data_tools.read_email_attachment(
        "org-1", "u-1", message_id="m1", filename="inv-2041.PDF"
    )
    assert result["filename"] == "INV-2041.pdf"
    assert "Invoice INV-2041" in result["text"]
    assert result["email"]["subject"] == "October invoice"

    missing = await data_tools.read_email_attachment("org-1", "u-1", message_id="m1", filename="nope.xlsx")
    assert "INV-2041.pdf" in missing["attachments"]


async def test_save_email_attachment_adds_it_to_files(monkeypatch) -> None:
    _connect_gmail(monkeypatch)
    _mock_http(monkeypatch)
    uploaded: dict = {}

    async def list_files(org_id, user_id):
        return [{"name": "INV-2041.pdf"}]

    async def upload(org_id, user_id, filename, data, name=None):
        uploaded.update(filename=filename, name=name, data=data)
        return {"file_id": "f9", "name": name, "kind": "pdf", "row_count": 1, "column_count": 0}

    monkeypatch.setattr(files_service, "list_files", list_files)
    monkeypatch.setattr(files_service, "upload_file", upload)
    state = ask_agent._FileToolState()

    result, *_ = await ask_agent._safe_run_tool(
        name="save_email_attachment", arguments={"message_id": "m1", "filename": "INV-2041.pdf"},
        org_id="org-1", user_id="u-1", people_cache={}, file_state=state,
    )
    assert result["saved"] is True
    assert uploaded["name"] == "INV-2041 (2).pdf"  # never overwrites an existing file
    assert uploaded["data"].startswith(b"%PDF")
    assert state.changes[0].kind == "pdf" and "October invoice" in state.changes[0].summary


async def test_gmail_sync_indexes_attachments_as_their_own_documents(monkeypatch) -> None:
    ingested: list = []
    downloads: list[str] = []

    async def fake_ingest(*, org_id, provider, external_id, version, conversation, document):
        ingested.append((conversation, document, provider, external_id, version))

    async def nothing_indexed(org_id, provider, external_id):
        return None

    def handler(request: httpx.Request) -> httpx.Response:
        downloads.append(request.url.path)
        return httpx.Response(200, json={"data": _b64(INVOICE)})

    monkeypatch.setattr(google_workspace, "ingest_external_source", fake_ingest)
    monkeypatch.setattr(google_workspace, "get_external_source", nothing_indexed)
    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
        count = await google_workspace._ingest_gmail_attachments(
            client, "token", MESSAGE, "Me@Acme.com", "org-1"
        )

    assert count == 1 and len(downloads) == 1
    conversation, document, provider, external_id, version = ingested[0]
    assert provider == "email_attachment" and version == "m1"
    assert external_id == "me@acme.com:m1:0:INV-2041.pdf"
    assert conversation.source == "email_attachment"
    assert "INV-2041.pdf" in conversation.title and "October invoice" in conversation.title
    assert all(m.text.startswith("[Attachment: INV-2041.pdf") for m in conversation.messages)
    assert "1,250.00" in conversation.messages[0].text
    assert document.visible_to == ["me@acme.com"]
    assert document.original_filename == "INV-2041.pdf"
    # The stored record names the mailbox, so the same file in two inboxes never
    # de-duplicates into one person's private copy.
    record = json.loads(document.data)
    assert record["account"] == "me@acme.com" and record["message_id"] == "m1"


def test_gmail_conversation_lists_attachments_and_reads_html_bodies() -> None:
    message = json.loads(json.dumps(MESSAGE))
    message["payload"]["parts"][0] = {
        "mimeType": "text/html", "filename": "",
        "body": {"data": _b64("<p>Hello <b>team</b></p><script>x()</script>")},
    }
    conversation = google_workspace._message_to_conversation(message, "me@acme.com")
    text = conversation.messages[0].text
    assert "Attachments: INV-2041.pdf, old.doc" in text
    assert "Hello team" in text and "<p>" not in text and "x()" not in text


@pytest.mark.parametrize("question", [
    "what's the total on the invoice Acme emailed me?",
    "summarize the PDF attached to Priya's latest email",
])
def test_attachment_questions_use_email_tools(question: str) -> None:
    assert ask_agent._wants_data(question) or ask_agent._wants_email(question)


async def test_gmail_resync_skips_attachments_already_indexed(monkeypatch) -> None:
    downloads: list[str] = []

    async def already(org_id, provider, external_id):
        return SimpleNamespace(status="active")

    async def never(**kwargs):
        raise AssertionError("an indexed attachment must not be ingested again")

    def handler(request: httpx.Request) -> httpx.Response:
        downloads.append(request.url.path)
        return httpx.Response(200, json={"data": _b64(INVOICE)})

    monkeypatch.setattr(google_workspace, "get_external_source", already)
    monkeypatch.setattr(google_workspace, "ingest_external_source", never)
    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
        count = await google_workspace._ingest_gmail_attachments(
            client, "token", MESSAGE, "me@acme.com", "org-1"
        )
    assert count == 0 and downloads == []
