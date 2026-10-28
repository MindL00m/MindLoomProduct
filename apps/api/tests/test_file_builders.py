"""Creating Excel, PDF and PowerPoint files from text."""

from __future__ import annotations

import io

import openpyxl
import pytest
from fastapi import HTTPException

from loombot import file_builders, files


def test_xlsx_from_markdown_tables_makes_one_sheet_per_table() -> None:
    data = file_builders.text_to_xlsx(
        "# Prices\n\n| Product | Price | Share |\n|---|---|---|\n| SOP | 1,250.50 | 12% |\n"
        "| MOP | 980 | 8% |\n| Total | =SUM(B2:B3) | |\n\n"
        "## Suppliers\n\n| Name | Country |\n|---|---|\n| Acme | India |\n"
    )
    workbook = openpyxl.load_workbook(io.BytesIO(data))
    assert workbook.sheetnames == ["Prices", "Suppliers"]
    prices = workbook["Prices"]
    assert prices["A1"].value == "Product" and prices["A1"].font.bold
    assert prices["B2"].value == 1250.5 and prices["B3"].value == 980
    assert prices["C2"].value == pytest.approx(0.12)
    assert prices["B4"].value == "=SUM(B2:B3)"
    assert prices.freeze_panes == "A2"


def test_xlsx_from_csv_text() -> None:
    data = file_builders.text_to_xlsx("item,qty\nbolts,40\nnuts,7\n")
    sheet = openpyxl.load_workbook(io.BytesIO(data)).active
    assert [[c.value for c in row] for row in sheet.iter_rows()] == [
        ["item", "qty"], ["bolts", 40], ["nuts", 7],
    ]


def test_pdf_from_markdown_has_the_text() -> None:
    data = file_builders.markdown_to_pdf(
        "# Quotation\n\nFor **Acme Ltd**.\n\n- Item one\n- Item two\n\n"
        "| Item | Qty |\n|---|---|\n| SOP | 20 MT |\n",
        title="Quotation",
    )
    assert data.startswith(b"%PDF")
    text = files.pdf_text(data)
    assert "Quotation" in text and "Acme Ltd" in text and "20 MT" in text


def test_long_pdf_flows_onto_more_pages() -> None:
    data = file_builders.markdown_to_pdf("\n\n".join(f"Paragraph {n} " * 20 for n in range(120)))
    assert files.binary_stats("pdf", data)[0] > 1


def test_pptx_from_markdown_one_slide_per_heading() -> None:
    data = file_builders.markdown_to_pptx(
        "# Q3 review\nAcme trading\n\n## Highlights\n- Revenue up\n  - mostly SOP\n- Costs flat\n\n"
        "## Prices\n| Product | Price |\n|---|---|\n| SOP | 520 |\n"
    )
    files.validate_binary("pptx", data)
    slides = file_builders.pptx_slides(data)
    assert [s["title"] for s in slides] == ["Q3 review", "Highlights", "Prices"]
    assert "Revenue up" in slides[1]["lines"] and "mostly SOP" in slides[1]["lines"]
    assert any("SOP | 520" in line for line in slides[2]["lines"])
    assert files.binary_stats("pptx", data) == (3, 0)
    assert "[Slide 2] Highlights" in files.pptx_text(data)


def test_build_binary_for_every_built_kind() -> None:
    for kind in sorted(files.BUILT_KINDS):
        data = files.build_binary(kind, "# Title\n\nHello\n\n| a | b |\n|---|---|\n| 1 | 2 |\n", f"x.{kind}")
        files.validate_binary(kind, data)


def test_new_text_kinds() -> None:
    assert files.kind_from_name("notes.md") == "md"
    assert files.kind_from_name("data.json") == "json"
    assert files.normalize_text("json", '{"a": 1}')[0] == '{"a": 1}'
    with pytest.raises(HTTPException):
        files.normalize_text("json", "{not json")
