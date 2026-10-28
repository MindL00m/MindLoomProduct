"""Shipment documents: extraction, exact calculations, and filled layouts."""

from __future__ import annotations

import io
from decimal import Decimal

import docx
import pytest

from loombot import ask_agent, files, shipment_docs, templates

RECORD = {
    "seller": {"name": "Acme Fertilisers Pvt Ltd", "address": "Mumbai, India", "country": "India"},
    "buyer": {"name": "Gulf Agro LLC", "address": "Jebel Ali, Dubai", "country": "UAE"},
    "consignee": {"name": "Gulf Agro LLC", "address": "Jebel Ali, Dubai"},
    "invoice_number": "INV-2026-042", "contract_number": "SC-118", "contract_date": "2026-09-02",
    "bl_number": "MSCU4471", "bl_date": "2026-10-20", "vessel": "MSC Aurora", "voyage": "V.214",
    "port_of_loading": "Nhava Sheva", "port_of_discharge": "Jebel Ali",
    "country_of_origin": "India", "incoterm": "CIF Jebel Ali", "currency": "USD",
    "containers": ["MSCU1234567", "MSCU7654321"],
    "items": [
        {"description": "SOP 52% granular", "hs_code": "31043000", "quantity": "20.000",
         "unit": "MT", "unit_price": "520.50", "gross_weight": "20,080", "net_weight": "20,000",
         "packages": "400", "package_type": "50 kg bags"},
        {"description": "MOP 60%", "hs_code": "31042000", "quantity": "10", "unit": "MT",
         "unit_price": "300", "amount": "3000", "gross_weight": "10,040", "net_weight": "10,000",
         "packages": "200", "package_type": "50 kg bags"},
    ],
    "analysis": [{"parameter": "K2O", "specification": "min 52%", "result": "52.3%", "method": "ISO"}],
    "missing": ["notify_party"],
}


def test_numbers_and_words() -> None:
    assert shipment_docs.to_decimal("20,000.500 MT") == Decimal("20000.500")
    assert shipment_docs.to_decimal("n/a") is None
    assert shipment_docs.amount_in_words(Decimal("13410.00"), "USD") == (
        "US dollars thirteen thousand four hundred and ten only"
    )
    assert shipment_docs.amount_in_words(Decimal("1005.5"), "EUR") == (
        "Euros one thousand and five and cents fifty only"
    )


def test_calculate_line_amounts_and_totals() -> None:
    figures, warnings = shipment_docs.calculate(RECORD)
    assert [i["amount"] for i in figures["items"]] == ["10,410.00", "3,000.00"]
    assert figures["total_amount"] == "13,410.00"
    assert figures["total_quantity"] == "30 MT"
    assert figures["total_gross_weight"] == "30,120" and figures["total_packages"] == "600"
    assert figures["amount_in_words"].startswith("US dollars thirteen thousand")
    assert warnings == []


def test_a_wrong_stated_amount_is_flagged_not_replaced() -> None:
    record = {"items": [{"quantity": "10", "unit_price": "300", "amount": "3100"}]}
    figures, warnings = shipment_docs.calculate(record)
    assert figures["items"][0]["amount"] == "3,100.00"
    assert "not quantity × price" in warnings[0]


@pytest.mark.parametrize("document_type", sorted(shipment_docs.DOCUMENT_TYPES))
def test_built_in_layouts_fill_completely(document_type) -> None:
    figures, _ = shipment_docs.calculate(RECORD)
    values = shipment_docs.flatten(RECORD, figures)
    out = templates.fill("docx", shipment_docs.built_in_docx(document_type), "", values)
    document = docx.Document(io.BytesIO(out))
    text = "\n".join(p.text for p in document.paragraphs) + "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    assert "Acme Fertilisers Pvt Ltd" in text
    # Only fields the shipment record does not have may be left over.
    leftover = [f for f in templates.unfilled("docx", out) if values.get(f)]
    assert leftover == [], leftover
    if document_type == "commercial_invoice":
        assert "13,410.00" in text and "MSCU1234567, MSCU7654321" in text
    if document_type == "certificate_of_analysis":
        assert "52.3%" in text and "SOP 52% granular; MOP 60%" in text


async def test_create_shipment_documents_uses_built_in_layouts(monkeypatch) -> None:
    async def extract(org_id, user_id, source_file_ids, instructions):
        assert source_file_ids == ["bl", "contract"]
        return RECORD

    saved: list = []

    async def fake_binary(org_id, user_id, name, kind, data, *, edited_by="user", **kwargs):
        saved.append((name, data))
        return {"file_id": str(len(saved)), "name": name, "kind": kind, "content": ""}

    async def no_files(org_id, user_id):
        return []

    monkeypatch.setattr(shipment_docs, "extract_shipment", extract)
    monkeypatch.setattr(files, "create_binary_file", fake_binary)
    monkeypatch.setattr(files, "list_files", no_files)
    result = await shipment_docs.create_shipment_documents(
        "o", "u", document_types=["commercial_invoice", "certificate_of_origin", "bogus"],
        source_file_ids=["bl", "contract"],
    )
    assert [d["file"]["name"] for d in result["documents"]] == [
        "Commercial invoice - INV-2026-042.docx", "Certificate of origin - INV-2026-042.docx",
    ]
    # Missing values show as blanks in the document but are still reported.
    invoice_text = "\n".join(p.text for p in docx.Document(io.BytesIO(saved[0][1])).paragraphs)
    assert "{{" not in invoice_text
    assert "notify_party_name" in result["documents"][0]["unfilled"]
    assert result["shipment"]["total_amount"] == "13,410.00"


async def test_no_valid_document_type_is_refused() -> None:
    from fastapi import HTTPException

    with pytest.raises(HTTPException):
        await shipment_docs.create_shipment_documents(
            "o", "u", document_types=["nope"], source_file_ids=[]
        )


async def test_ask_tool_needs_no_file_name(monkeypatch) -> None:
    async def fake_create(org_id, user_id, **kwargs):
        return {"documents": [], "shipment": {"total_amount": "1.00"}, "missing": [], "warnings": []}

    monkeypatch.setattr(shipment_docs, "create_shipment_documents", fake_create)
    out = await ask_agent._run_file_tool(
        "create_shipment_documents", {"document_types": ["commercial_invoice"]},
        org_id="o", user_id="u", state=ask_agent._FileToolState(),
    )
    assert out["totals"] == {"total_amount": "1.00"}
