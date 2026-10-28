"""Templates: placeholders, line items, filling, and samples made into templates."""

from __future__ import annotations

import io

import docx
import openpyxl
import pytest

from loombot import file_builders, files, templates

INVOICE_MD = """# Commercial invoice {{invoice_number}}

Buyer: **{{buyer_name}}**, {{buyer_address}}

Date: {{invoice_date}}

| Description | Qty (MT) | Rate (USD) | Amount (USD) |
|---|---|---|---|
| {{items.description}} | {{items.quantity}} | {{items.rate}} | {{items.amount}} |

Total: {{total_amount}}
"""

ITEMS = [
    {"description": "SOP 52%", "quantity": 20, "rate": 520, "amount": 10400},
    {"description": "MOP 60%", "quantity": 10, "rate": 300, "amount": 3000},
]


def _docx_template() -> bytes:
    from loombot.docx_builder import markdown_to_docx

    return markdown_to_docx(INVOICE_MD, title="Invoice")


def test_fields_and_line_items_are_found() -> None:
    found = templates.template_fields("docx", _docx_template())
    assert found["fields"] == [
        "invoice_number", "buyer_name", "buyer_address", "invoice_date", "total_amount",
    ]
    assert found["lists"] == {"items": ["description", "quantity", "rate", "amount"]}


def test_docx_fill_repeats_item_rows_and_keeps_bold() -> None:
    values = {
        "invoice_number": "INV-042", "buyer_name": "Gulf Agro LLC", "buyer_address": "Dubai",
        "invoice_date": "2026-10-28", "total_amount": "13,400.00", "items": ITEMS,
    }
    out = templates.fill("docx", _docx_template(), "", values)
    document = docx.Document(io.BytesIO(out))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Commercial invoice INV-042" in text and "Total: 13,400.00" in text
    assert any(run.bold and run.text == "Gulf Agro LLC" for p in document.paragraphs for run in p.runs)
    rows = [[c.text for c in row.cells] for row in document.tables[0].rows]
    assert rows[1:] == [["SOP 52%", "20", "520", "10400"], ["MOP 60%", "10", "300", "3000"]]
    assert templates.unfilled("docx", out) == []


def test_unfilled_fields_are_reported() -> None:
    out = templates.fill("docx", _docx_template(), "", {"invoice_number": "1"})
    assert "buyer_name" in templates.unfilled("docx", out)


def test_xlsx_fill_inserts_rows_and_grows_totals() -> None:
    data = file_builders.rows_to_xlsx([("Invoice", [
        ["Invoice", "{{invoice_number}}", "", ""],
        ["Description", "Qty", "Rate", "Amount"],
        ["{{items.description}}", "{{items.quantity}}", "{{items.rate}}", "=B3*C3"],
        ["Total", "", "", "=SUM(D3:D3)"],
    ])])
    out = templates.fill("xlsx", data, "", {"invoice_number": "INV-7", "items": ITEMS})
    sheet = openpyxl.load_workbook(io.BytesIO(out)).active
    rows = [[c.value for c in row] for row in sheet.iter_rows()]
    assert rows[0][1] == "INV-7"
    assert rows[2] == ["SOP 52%", 20, 520, "=B3*C3"]
    assert rows[3] == ["MOP 60%", 10, 300, "=B4*C4"]
    assert rows[4] == ["Total", None, None, "=SUM(D3:D4)"]


def test_csv_and_text_fill() -> None:
    csv_out = templates.fill(
        "csv", b"", "item,qty\n{{items.description}},{{items.quantity}}\n", {"items": ITEMS}
    )
    assert csv_out == "item,qty\nSOP 52%,20\nMOP 60%,10\n"
    text_out = templates.fill("txt", b"", "Dear {{name}},\n- {{items.description}}\n", {
        "name": "Priya", "items": ITEMS,
    })
    assert text_out == "Dear Priya,\n- SOP 52%\n- MOP 60%\n"


def test_pptx_fill() -> None:
    data = file_builders.markdown_to_pptx("# Offer for {{buyer}}\n\n## Prices\n- {{items.description}}: {{items.rate}}\n")
    out = templates.fill("pptx", data, "", {"buyer": "Acme", "items": ITEMS})
    slides = file_builders.pptx_slides(out)
    assert slides[0]["title"] == "Offer for Acme"
    assert slides[1]["lines"] == ["SOP 52%: 520; MOP 60%: 300"]


async def test_fill_template_file_uses_loombot_for_missing_values(monkeypatch) -> None:
    data = _docx_template()

    async def fake_bytes(org_id, user_id, file_id):
        return data, {"name": "Invoice (template).docx", "kind": "docx"}

    async def fake_values(org_id, user_id, **kwargs):
        assert kwargs["source_file_ids"] == ["bl"]
        return {"values": {"buyer_name": "From BL", "items": ITEMS}, "label": "Gulf Agro INV-9",
                "missing": [], "notes": "From the bill of lading."}

    saved: dict = {}

    async def fake_save(org_id, user_id, meta, output, *, name, label, based_on, edited_by="loombot"):
        saved.update(label=label, output=output, based_on=based_on)
        return {"file_id": "new", "name": f"Invoice - {label}.docx", "kind": "docx"}

    monkeypatch.setattr(files, "get_file_bytes", fake_bytes)
    monkeypatch.setattr(templates, "values_from_sources", fake_values)
    monkeypatch.setattr(templates, "_save_output", fake_save)
    result = await templates.fill_template_file(
        "o", "u", "tpl", values={"invoice_number": "INV-9"}, instructions="invoice for Gulf Agro",
        source_file_ids=["bl"],
    )
    assert result["values"]["invoice_number"] == "INV-9"  # typed value wins
    assert result["values"]["buyer_name"] == "From BL"
    assert saved["label"] == "Gulf Agro INV-9" and saved["based_on"] == "tpl"
    assert "invoice_date" in result["unfilled"]


async def test_files_without_placeholders_cannot_be_filled(monkeypatch) -> None:
    from fastapi import HTTPException
    from loombot.docx_builder import markdown_to_docx

    async def fake_bytes(org_id, user_id, file_id):
        return markdown_to_docx("Plain sample"), {"name": "Sample.docx", "kind": "docx"}

    monkeypatch.setattr(files, "get_file_bytes", fake_bytes)
    with pytest.raises(HTTPException) as exc:
        await templates.fill_template_file("o", "u", "f")
    assert "Make template" in exc.value.detail


async def test_make_template_from_sample(monkeypatch) -> None:
    from loombot.docx_builder import markdown_to_docx

    sample = markdown_to_docx(
        "# Invoice 0017\n\nBuyer: Gulf Agro LLC\n\n| Item | Qty |\n|---|---|\n| SOP | 20 |\n| MOP | 10 |\n"
    )

    async def fake_bytes(org_id, user_id, file_id):
        return sample, {"name": "Invoice 0017.docx", "kind": "docx"}

    async def fake_chat(org_id, system, user):
        assert "Segments" in user
        return {"edits": [
            {"id": "p0", "text": "Invoice {{invoice_number}}"},
            {"id": "p1", "text": "Buyer: {{buyer_name}}"},
            {"id": "p4", "text": "{{items.item}}"}, {"id": "p5", "text": "{{items.qty}}"},
            {"id": "p6", "text": ""}, {"id": "p7", "text": ""},
        ], "summary": "Made fields."}

    saved: dict = {}

    async def fake_save(org_id, user_id, meta, output, *, name, label, based_on, edited_by="loombot"):
        saved.update(name=name, output=output)
        return {"file_id": "t", "name": name, "kind": "docx"}

    async def fake_flag(org_id, user_id, file_id, flag):
        saved["flag"] = flag
        return {}

    monkeypatch.setattr(files, "get_file_bytes", fake_bytes)
    monkeypatch.setattr("lens.status_trackers._chat_json", fake_chat)
    monkeypatch.setattr(templates, "_save_output", fake_save)
    monkeypatch.setattr(files, "set_template", fake_flag)
    result = await templates.make_template_from_sample("o", "u", "f")
    assert saved["name"] == "Invoice 0017 (template).docx" and saved["flag"] is True
    assert result["fields"] == ["invoice_number", "buyer_name"]
    assert result["lists"] == {"items": ["item", "qty"]}
    table = docx.Document(io.BytesIO(saved["output"])).tables[0]
    assert len(table.rows) == 2  # the emptied second item row was removed
