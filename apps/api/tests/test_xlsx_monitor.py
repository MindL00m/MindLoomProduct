"""Excel workbooks keep themselves updated like CSV files."""

from __future__ import annotations

import io

import openpyxl
from openpyxl.styles import Font

from loombot import file_monitor


def _price_sheet() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Prices"
    sheet["A1"] = "Supplier price list"  # a title above the table
    sheet.append([])
    sheet.append(["Supplier", "Product", "Price", "Qty", "Value"])
    sheet.append(["Acme", "SOP", 520, 20, "=C4*D4"])
    sheet["C4"].font = Font(bold=True)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_matrix_skips_a_title_above_the_table() -> None:
    matrix = file_monitor.xlsx_matrix(_price_sheet())
    assert matrix[0] == ["Supplier", "Product", "Price", "Qty", "Value"]
    assert matrix[1][:4] == ["Acme", "SOP", "520", "20"]


def test_rows_are_appended_below_a_titled_table() -> None:
    out, added, _ = file_monitor.apply_xlsx_changes(_price_sheet(), [["B", "MOP", "300", "1", ""]], [])
    sheet = openpyxl.load_workbook(io.BytesIO(out)).active
    assert added == 1 and [c.value for c in sheet[5]] == ["B", "MOP", 300, 1, "=C5*D5"]


def _table() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Supplier", "Product", "Price", "Qty", "Value"])
    sheet.append(["Acme", "SOP", 520, 20, "=C2*D2"])
    sheet["C2"].font = Font(bold=True)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_new_rows_keep_style_numbers_and_row_formulas() -> None:
    out, added, updated = file_monitor.apply_xlsx_changes(
        _table(),
        [["Gulf Agro", "MOP", "300.50", "10", ""], ["Acme", "SOP", "520", "20", "=C2*D2"]],
        [],
    )
    assert (added, updated) == (1, 0)  # the second proposed row is already there
    sheet = openpyxl.load_workbook(io.BytesIO(out)).active
    assert [c.value for c in sheet[3]] == ["Gulf Agro", "MOP", 300.5, 10, "=C3*D3"]
    assert sheet["C3"].font.bold


def test_price_update_changes_value_but_never_a_formula() -> None:
    out, added, updated = file_monitor.apply_xlsx_changes(
        _table(),
        [],
        [{"row": 1, "column": "Price", "value": "535"}, {"row": 1, "column": "Value", "value": "1"}],
    )
    assert (added, updated) == (0, 1)
    sheet = openpyxl.load_workbook(io.BytesIO(out)).active
    assert sheet["C2"].value == 535 and sheet["E2"].value == "=C2*D2"


def test_nothing_to_change_returns_the_same_bytes() -> None:
    data = _table()
    assert file_monitor.apply_xlsx_changes(data, [], []) == (data, 0, 0)


async def test_monitor_prompt_shows_workbook_rows(monkeypatch) -> None:
    async def fake_get(self, key):
        return _table()

    monkeypatch.setattr("core.blob_storage.LocalBlobStorage.get", fake_get, raising=False)
    from core import blob_storage

    class Storage:
        async def get(self, key):
            return _table()

    monkeypatch.setattr(blob_storage, "get_blob_storage", lambda: Storage())
    matrix = await file_monitor._file_matrix({"kind": "xlsx", "blob_key": "k"})
    prompt = file_monitor._monitor_prompt(
        {"name": "prices.xlsx", "monitor_instructions": "Update prices", "content": ""},
        {"source_label": "Email from Acme"}, [], matrix,
    )
    assert "Header: Supplier,Product,Price,Qty,Value" in prompt
    assert "1: Acme,SOP,520,20," in prompt
